from django.shortcuts import render
from django.views import View
from django.contrib.auth.mixins import LoginRequiredMixin
from django.utils import timezone
from django.db.models import Sum, Count, Q
from decimal import Decimal
from datetime import datetime

from apps.sales.models import Sale, SaleItem
from apps.customers.models import CustomerTransaction
from apps.accounting.models import CashTransfer, BankTransfer, ExpenditureItem
from apps.payments.models import ECashLedger

class EndOfDaySummaryView(LoginRequiredMixin, View):
    """
    Role-specific end of day summary report.
    Shows the daily activity and financial summary for the requested date.
    """
    template_name = 'reports/end_of_day_summary.html'

    def get(self, request):
        user = request.user
        tenant = user.tenant
        role_name = user.role.name if user.role else None

        # Determine target dates
        start_date_str = request.GET.get('start_date')
        end_date_str = request.GET.get('end_date')
        
        today = timezone.now().date()
        
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date() if start_date_str else today
        except ValueError:
            start_date = today
            
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date() if end_date_str else today
        except ValueError:
            end_date = today

        # Ensure start_date is before or equal to end_date
        if start_date > end_date:
            start_date = end_date

        context = {
            'start_date': start_date,
            'end_date': end_date,
            'role_name': role_name,
            'role_display_name': role_name.replace('_', ' ').title() if role_name else 'Unknown',
        }

        # Shared base filters
        date_filter = {
            'created_at__date__gte': start_date,
            'created_at__date__lte': end_date,
            'tenant': tenant
        }

        # If Accountant/Admin is viewing a specific shop
        viewing_shop_id = request.GET.get('shop_id')
        viewing_shop = None
        
        if role_name in ['ACCOUNTANT', 'ADMIN', 'AUDITOR']:
            from apps.core.models import Location
            context['available_shops'] = Location.objects.filter(tenant=tenant, location_type='SHOP', is_active=True)
            
            if viewing_shop_id:
                try:
                    viewing_shop = Location.objects.get(id=viewing_shop_id, tenant=tenant, location_type='SHOP')
                    context['viewing_shop'] = viewing_shop
                except Location.DoesNotExist:
                    pass

        from django.core.cache import cache
        cache_key = f'eod_summary_{user.id}_{start_date}_{end_date}_{viewing_shop_id}'
        cached_data = cache.get(cache_key)
        if cached_data:
            context.update(cached_data)
        else:
            summary_context = {}
            # ---------------------------------------------------------
            # ATTENDANT SUMMARY
            # ---------------------------------------------------------
            if role_name == 'SHOP_ATTENDANT':
                sales = Sale.objects.filter(**date_filter, attendant=user).exclude(status='PENDING')
                summary_context['total_sales_value'] = sales.aggregate(t=Sum('total'))['t'] or Decimal('0.00')
                summary_context['invoice_count'] = sales.count()
                
                summary_context['items_sold_qty'] = SaleItem.objects.filter(
                    sale__in=sales
                ).aggregate(q=Sum('quantity'))['q'] or Decimal('0')
    
                summary_context['customer_payments_received'] = CustomerTransaction.objects.filter(
                    **date_filter,
                    performed_by=user,
                    transaction_type='CREDIT'
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
            # ---------------------------------------------------------
            # CASHIER SUMMARY
            # ---------------------------------------------------------
            elif role_name == 'SHOP_CASHIER':
                sales = Sale.objects.filter(**date_filter, cashier=user).exclude(status='PENDING')
                summary_context['invoice_count'] = sales.count()
                
                methods = sales.values('payment_method').annotate(total=Sum('amount_paid'))
                payments_breakdown = {m['payment_method']: m['total'] for m in methods}
                summary_context['payments_breakdown'] = payments_breakdown
                summary_context['total_sales_value'] = sum(payments_breakdown.values()) or Decimal('0.00')
    
                summary_context['customer_payments_received'] = CustomerTransaction.objects.filter(
                    **date_filter,
                    performed_by=user,
                    transaction_type='CREDIT'
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
                summary_context['bank_transfers_made'] = BankTransfer.objects.filter(
                    **date_filter,
                    accountant=user
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
            # ---------------------------------------------------------
            # SHOP MANAGER OR ACCOUNTANT VIEWING SHOP SUMMARY
            # ---------------------------------------------------------
            elif role_name == 'SHOP_MANAGER' or viewing_shop:
                shop = viewing_shop if viewing_shop else user.location
                if shop:
                    sales = Sale.objects.filter(**date_filter, shop=shop).exclude(status='PENDING')
                    summary_context['shop_total_sales_value'] = sales.aggregate(t=Sum('total'))['t'] or Decimal('0.00')
                    summary_context['shop_invoice_count'] = sales.count()
                    
                    summary_context['shop_items_sold_qty'] = SaleItem.objects.filter(
                        sale__in=sales
                    ).aggregate(q=Sum('quantity'))['q'] or Decimal('0')
    
                    summary_context['shop_customer_payments_received'] = CustomerTransaction.objects.filter(
                        **date_filter,
                        customer__shop=shop,
                        transaction_type='CREDIT'
                    ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
                    methods = sales.values('payment_method').annotate(total=Sum('amount_paid'))
                    payments_breakdown = {m['payment_method']: m['total'] for m in methods if m['total'] and m['total'] > 0}
                    
                    # Debt generated (unpaid amounts from sales)
                    total_billed = summary_context['shop_total_sales_value']
                    total_paid = sales.aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
                    credit_sales_generated = total_billed - total_paid
                    if credit_sales_generated > 0:
                        payments_breakdown['UNPAID (CREDIT SALES)'] = credit_sales_generated
    
                    summary_context['payments_breakdown'] = payments_breakdown
    
                    summary_context['transferred_funds_out'] = CashTransfer.objects.filter(
                        tenant=tenant,
                        from_location=shop,
                        status='CONFIRMED',
                        transfer_type='DEPOSIT',
                        confirmed_at__date__gte=start_date,
                        confirmed_at__date__lte=end_date
                    ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
                    summary_context['shop_expenditures'] = ExpenditureItem.objects.filter(
                        request__tenant=tenant,
                        request__location=shop,
                        status='APPROVED',
                        approved_at__date__gte=start_date,
                        approved_at__date__lte=end_date
                    ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                    
                    # Internal funds received - skip this if Accountant viewing the shop, or if using strict workflow
                    if role_name == 'SHOP_MANAGER' and not tenant.use_strict_sales_workflow:
                        summary_context['transferred_funds_in'] = CashTransfer.objects.filter(
                            tenant=tenant,
                            to_user=user,
                            status='CONFIRMED',
                            confirmed_at__date__gte=start_date,
                            confirmed_at__date__lte=end_date
                        ).exclude(from_user=user).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                    
                    # Bank transfers made by cashiers in this shop (only in strict workflow)
                    if tenant.use_strict_sales_workflow:
                        summary_context['shop_bank_transfers_made'] = BankTransfer.objects.filter(
                            tenant=tenant,
                            accountant__location=shop,
                            created_at__date__gte=start_date,
                            created_at__date__lte=end_date
                        ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
            # ---------------------------------------------------------
            # OVERALL VIEW FOR ACCOUNTANT, ADMIN, AUDITOR
            # ---------------------------------------------------------
            elif role_name in ['ACCOUNTANT', 'ADMIN', 'AUDITOR']:
                sales = Sale.objects.filter(**date_filter).exclude(status='PENDING')
                methods = sales.values('payment_method').annotate(total=Sum('amount_paid'))
                payments_breakdown = {m['payment_method']: m['total'] for m in methods if m['total'] and m['total'] > 0}
                
                total_billed = sales.aggregate(t=Sum('total'))['t'] or Decimal('0.00')
                total_paid = sales.aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
                credit_sales_generated = total_billed - total_paid
                if credit_sales_generated > 0:
                    payments_breakdown['UNPAID (CREDIT SALES)'] = credit_sales_generated
    
                summary_context['payments_breakdown'] = payments_breakdown
    
                summary_context['cash_transfers_received'] = CashTransfer.objects.filter(
                    tenant=tenant,
                    to_user=user,
                    status='CONFIRMED',
                    transfer_type='DEPOSIT',
                    confirmed_at__date__gte=start_date,
                    confirmed_at__date__lte=end_date
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
                # Bank transfers made by THIS user only
                summary_context['bank_transfers_made'] = BankTransfer.objects.filter(
                    **date_filter,
                    accountant=user
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
                summary_context['expenditures_disbursed'] = ExpenditureItem.objects.filter(
                    request__tenant=tenant,
                    status='APPROVED',
                    approved_at__date__gte=start_date,
                    approved_at__date__lte=end_date,
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
                summary_context['digital_funds_withdrawn'] = ECashLedger.objects.filter(
                    **date_filter,
                    transaction_type='WITHDRAWAL',
                    created_by=user
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                
                if summary_context['digital_funds_withdrawn'] < 0:
                    summary_context['digital_funds_withdrawn'] = abs(summary_context['digital_funds_withdrawn'])
                    
            cache.set(cache_key, summary_context, 600)  # 10 minutes
            context.update(summary_context)
        if request.GET.get('pdf') == 'true':
            from django.template.loader import get_template
            from django.http import HttpResponse
            from xhtml2pdf import pisa
            from io import BytesIO
            
            template = get_template('reports/pdf_summary_template.html')
            html = template.render(context)
            result = BytesIO()
            pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
            if not pdf.err:
                response = HttpResponse(result.getvalue(), content_type='application/pdf')
                response['Content-Disposition'] = f'attachment; filename="end_of_day_summary_{start_date}.pdf"'
                return response

        return render(request, self.template_name, context)

from apps.inventory.models import InventoryLedger

class EndOfDayDetailsView(LoginRequiredMixin, View):
    template_name = 'reports/end_of_day_details.html'

    def get(self, request):
        user = request.user
        tenant = user.tenant
        role_name = user.role.name if user.role else None

        start_date_str = request.GET.get('start_date')
        end_date_str = request.GET.get('end_date')
        
        today = timezone.now().date()
        
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date() if start_date_str else today
        except ValueError:
            start_date = today
            
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date() if end_date_str else today
        except ValueError:
            end_date = today

        if start_date > end_date:
            start_date = end_date
            
        exclude_inactive = request.GET.get('exclude_inactive') == 'true'
        top_n_str = request.GET.get('top_n', 'all')
        
        viewing_shop_id = request.GET.get('shop_id')
        viewing_shop = None
        
        if role_name in ['ACCOUNTANT', 'ADMIN', 'AUDITOR']:
            from apps.core.models import Location
            if viewing_shop_id:
                try:
                    viewing_shop = Location.objects.get(id=viewing_shop_id, tenant=tenant, location_type='SHOP')
                except Location.DoesNotExist:
                    pass

        from django.core.cache import cache
        cache_key = f'eod_details_{user.id}_{start_date}_{end_date}_{viewing_shop_id}_{exclude_inactive}_{top_n_str}'
        cached_data = cache.get(cache_key)
        if cached_data:
            sorted_stats, summary = cached_data
        else:
            sales = Sale.objects.filter(
                tenant=tenant,
                created_at__date__gte=start_date,
                created_at__date__lte=end_date,
                status__in=['COMPLETED', 'PENDING_DISPATCH']
            )
            
            shop_filter_for_inventory = None
    
            if role_name == 'SHOP_ATTENDANT':
                sales = sales.filter(attendant=user)
                shop_filter_for_inventory = user.location
            elif role_name == 'SHOP_CASHIER':
                sales = sales.filter(cashier=user)
                shop_filter_for_inventory = user.location
            elif role_name == 'SHOP_MANAGER' or viewing_shop:
                shop = viewing_shop if viewing_shop else user.location
                sales = sales.filter(shop=shop)
                shop_filter_for_inventory = shop
    
            items_agg = SaleItem.objects.filter(sale__in=sales).values(
                'product_id', 'product__name'
            ).annotate(
                total_qty=Sum('quantity'),
                total_revenue=Sum('total')
            ).order_by('-total_qty')
            
            product_stats = {}
            for item in items_agg:
                product_stats[item['product_id']] = {
                    'name': item['product__name'],
                    'qty_sold': item['total_qty'],
                    'revenue': item['total_revenue'],
                    'stock_left': 'N/A'
                }
                
            if shop_filter_for_inventory:
                if not exclude_inactive:
                    # Pull all products at this location from the ledger
                    ledger_agg = InventoryLedger.objects.filter(
                        tenant=tenant, location=shop_filter_for_inventory
                    ).values('product_id', 'product__name').annotate(
                        stock=Sum('quantity')
                    )
                    for row in ledger_agg:
                        pid = row['product_id']
                        stock = row['stock'] or Decimal('0')
                        if pid not in product_stats:
                            product_stats[pid] = {
                                'name': row['product__name'],
                                'qty_sold': Decimal('0'),
                                'revenue': Decimal('0.00'),
                                'stock_left': stock
                            }
                        else:
                            product_stats[pid]['stock_left'] = stock
                else:
                    # Only look up stock for products that had sales
                    if product_stats:
                        ledger_agg = InventoryLedger.objects.filter(
                            tenant=tenant,
                            location=shop_filter_for_inventory,
                            product_id__in=product_stats.keys()
                        ).values('product_id').annotate(stock=Sum('quantity'))
                        for row in ledger_agg:
                            product_stats[row['product_id']]['stock_left'] = row['stock'] or Decimal('0')
                        
            sorted_stats = sorted(product_stats.values(), key=lambda x: x['qty_sold'], reverse=True)
            
            others_qty = Decimal('0')
            others_rev = Decimal('0.00')
            
            if top_n_str in ['10', '20']:
                n = int(top_n_str)
                if len(sorted_stats) > n:
                    top_items = sorted_stats[:n]
                    others = sorted_stats[n:]
                    for o in others:
                        others_qty += o['qty_sold']
                        others_rev += o['revenue']
                        
                    if others_qty > 0 or others_rev > 0:
                        top_items.append({
                            'name': 'Others',
                            'qty_sold': others_qty,
                            'revenue': others_rev,
                            'stock_left': '-'
                        })
                    sorted_stats = top_items
    
            date_filter = {
                'created_at__date__gte': start_date,
                'created_at__date__lte': end_date,
                'tenant': tenant
            }
            summary = {}
            if role_name == 'SHOP_ATTENDANT':
                summ_sales = Sale.objects.filter(**date_filter, attendant=user).exclude(status='PENDING')
                summary['total_sales_value'] = summ_sales.aggregate(t=Sum('total'))['t'] or Decimal('0.00')
                summary['invoice_count'] = summ_sales.count()
                summary['items_sold_qty'] = SaleItem.objects.filter(sale__in=summ_sales).aggregate(q=Sum('quantity'))['q'] or Decimal('0')
                summary['customer_payments_received'] = CustomerTransaction.objects.filter(
                    **date_filter, performed_by=user, transaction_type='CREDIT'
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
            elif role_name == 'SHOP_CASHIER':
                summ_sales = Sale.objects.filter(**date_filter, cashier=user).exclude(status='PENDING')
                summary['invoice_count'] = summ_sales.count()
                methods = summ_sales.values('payment_method').annotate(total=Sum('amount_paid'))
                payments_breakdown = {m['payment_method']: m['total'] for m in methods}
                summary['payments_breakdown'] = payments_breakdown
                summary['total_sales_value'] = sum(payments_breakdown.values()) or Decimal('0.00')
                summary['customer_payments_received'] = CustomerTransaction.objects.filter(
                    **date_filter, performed_by=user, transaction_type='CREDIT'
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                summary['bank_transfers_made'] = BankTransfer.objects.filter(
                    **date_filter, accountant=user
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
            elif role_name == 'SHOP_MANAGER' or viewing_shop:
                shop = viewing_shop if viewing_shop else user.location
                if shop:
                    summ_sales = Sale.objects.filter(**date_filter, shop=shop).exclude(status='PENDING')
                    summary['shop_total_sales_value'] = summ_sales.aggregate(t=Sum('total'))['t'] or Decimal('0.00')
                    summary['shop_invoice_count'] = summ_sales.count()
                    summary['shop_items_sold_qty'] = SaleItem.objects.filter(sale__in=summ_sales).aggregate(q=Sum('quantity'))['q'] or Decimal('0')
                    summary['shop_customer_payments_received'] = CustomerTransaction.objects.filter(
                        **date_filter, customer__shop=shop, transaction_type='CREDIT'
                    ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                    
                    methods = summ_sales.values('payment_method').annotate(total=Sum('amount_paid'))
                    payments_breakdown = {m['payment_method']: m['total'] for m in methods if m['total'] and m['total'] > 0}
                    
                    total_billed = summary['shop_total_sales_value']
                    total_paid = summ_sales.aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
                    credit_sales_generated = total_billed - total_paid
                    if credit_sales_generated > 0:
                        payments_breakdown['UNPAID (CREDIT SALES)'] = credit_sales_generated
                        
                    summary['payments_breakdown'] = payments_breakdown
    
                    summary['transferred_funds_out'] = CashTransfer.objects.filter(
                        tenant=tenant, from_location=shop, status='CONFIRMED', transfer_type='DEPOSIT',
                        confirmed_at__date__gte=start_date, confirmed_at__date__lte=end_date
                    ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                    summary['shop_expenditures'] = ExpenditureItem.objects.filter(
                        request__tenant=tenant, request__location=shop, status='APPROVED',
                        approved_at__date__gte=start_date, approved_at__date__lte=end_date
                    ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                    summary['shop_bank_transfers_made'] = BankTransfer.objects.filter(
                        tenant=tenant, accountant__location=shop,
                        created_at__date__gte=start_date, created_at__date__lte=end_date
                    ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
    
            elif role_name in ['ACCOUNTANT', 'ADMIN', 'AUDITOR'] and not viewing_shop:
                summary['cash_transfers_received'] = CashTransfer.objects.filter(
                    tenant=tenant, to_user=user, status='CONFIRMED', transfer_type='DEPOSIT',
                    confirmed_at__date__gte=start_date, confirmed_at__date__lte=end_date
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                
                methods = sales.values('payment_method').annotate(total=Sum('amount_paid'))
                payments_breakdown = {m['payment_method']: m['total'] for m in methods if m['total'] and m['total'] > 0}
                
                total_billed = sales.aggregate(t=Sum('total'))['t'] or Decimal('0.00')
                total_paid = sales.aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
                credit_sales_generated = total_billed - total_paid
                if credit_sales_generated > 0:
                    payments_breakdown['UNPAID (CREDIT SALES)'] = credit_sales_generated
                    
                summary['payments_breakdown'] = payments_breakdown
                
                summary['bank_transfers_made'] = BankTransfer.objects.filter(
                    **date_filter, accountant=user
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                summary['expenditures_disbursed'] = ExpenditureItem.objects.filter(
                    request__tenant=tenant, status='APPROVED',
                    approved_at__date__gte=start_date, approved_at__date__lte=end_date
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                digital = ECashLedger.objects.filter(
                    **date_filter, transaction_type='WITHDRAWAL', created_by=user
                ).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                summary['digital_funds_withdrawn'] = abs(digital)
    
            cache.set(cache_key, (sorted_stats, summary), 600)  # 10 minutes
        context = {
            'start_date': start_date,
            'end_date': end_date,
            'role_name': role_name,
            'viewing_shop': viewing_shop,
            'stats': sorted_stats,
            'exclude_inactive': exclude_inactive,
            'top_n': top_n_str,
            'shop_id': viewing_shop_id or '',
            **summary,
        }

        if request.GET.get('pdf') == 'true':
            from django.template.loader import get_template
            from django.http import HttpResponse
            from xhtml2pdf import pisa
            from io import BytesIO

            template = get_template('reports/pdf_combined_template.html')
            html = template.render(context)
            result = BytesIO()
            pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
            if not pdf.err:
                response = HttpResponse(result.getvalue(), content_type='application/pdf')
                response['Content-Disposition'] = f'attachment; filename="end_of_day_report_{start_date}.pdf"'
                return response

        if request.GET.get('excel') == 'true':
            import openpyxl
            from django.http import HttpResponse

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "End of Day Report"
            
            ws.append(["End of Day Report"])
            ws.append(["Period:", f"{start_date} to {end_date}"])
            ws.append(["Role:", role_name.replace('_', ' ').title()])
            if viewing_shop:
                ws.append(["Shop:", viewing_shop.name])
            ws.append([])
            
            ws.append(["Financial Summary"])
            for k, v in summary.items():
                if k == 'payments_breakdown':
                    continue
                label = k.replace('_', ' ').title()
                ws.append([label, float(v) if isinstance(v, Decimal) else v])
                
            if 'payments_breakdown' in summary and summary['payments_breakdown']:
                ws.append([])
                ws.append(["Payments Breakdown"])
                for method, amount in summary['payments_breakdown'].items():
                    ws.append([method.replace('_', ' ').title(), float(amount) if isinstance(amount, Decimal) else amount])
            
            ws.append([])
            ws.append(["Detailed Product Breakdown"])
            ws.append(['#', 'Product Name', 'Qty Sold', 'Total Revenue', 'Stock Left'])
            
            for i, item in enumerate(sorted_stats, 1):
                idx = i if item['name'] != 'Others' else '-'
                qty = float(item['qty_sold']) if item['qty_sold'] else 0
                rev = float(item['revenue']) if item['revenue'] else 0
                stock = item['stock_left']
                if isinstance(stock, Decimal):
                    stock = float(stock)
                ws.append([idx, item['name'], qty, rev, stock])
                
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="end_of_day_report_{start_date}.xlsx"'
            wb.save(response)
            return response

        return render(request, self.template_name, context)
